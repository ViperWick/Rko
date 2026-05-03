# -*- coding: utf-8 -*-
"""
產生「自學課程」雙欄版：PowerPoint + Word。

左欄：觀念／目標　右欄：操作步驟／指令／自評

執行（repo 根目錄）：
  py -m pip install -r docs/requirements_handouts.txt
  py docs/generate_course_handouts.py

輸出（本目錄 docs/）：
  自學課程_雙排介紹.pptx
  自學課程_雙排介紹.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt

DOCS = Path(__file__).resolve().parent
OUT_PPTX = DOCS / "自學課程_雙排介紹.pptx"
OUT_DOCX = DOCS / "自學課程_雙排介紹.docx"

# 每課：title, left（觀念）, right（操作／自評）
LESSONS: list[dict[str, str]] = [
    {
        "title": "封面",
        "left": "自學課程：AI × 指令 × 專案架構\n一步一步（雙欄速覽版）\n\n建議：每天 1 課，20～40 分鐘\n詳全文：docs/自學課程_AI與專案實操_一步一步.md",
        "right": "搭配閱讀：\ndocs/AI與程式入門_專案架構總覽.md\n\n專案根目錄範例：\nC:\\Users\\User\\.cursor-tutor\n\n完成一課可對 Cursor 說：\n「我做完第 N 課…」",
    },
    {
        "title": "架構一覽（入門總覽精簡）",
        "left": "• Cursor 里的 AI：幫你讀寫專案、給指令\n• 終端機 / .bat：真正執行程式\n• Ollama：本機大模型 + API（常見 11434）\n• 雲端（Claude/GPT/Codex）：要帳號，常計費\n• Agent：模型 + 你的程式（工具／迴圈）",
        "right": "repo 資料夾：\n• ai_agents → 統一雙擊入口\n• etoro_trading → eToro 主專案\n• local_agent_demo → 教學 Agent\n• docs → 說明與本課程",
    },
    {
        "title": "課程地圖（0～12）",
        "left": "0 專案地圖\n1 檔案總管與根目錄\n2 終端機 cd\n3 py 跑一行\n4 py -m 入口\n5 pip／套件≠付費\n6 Ollama 服務\n7 ai_agents 選單\n8 eToro 摘要與 log\n9 讀 config.py\n10 改 SYSTEM 提示\n11 .env 本機vs雲端\n12 排程心態",
        "right": "自評每課都要做。\n畢業後可重跑第 7 課起加深記憶。\n\n延伸：\netoro_trading/docs/簡易操作流程.md",
    },
    {
        "title": "第 0 課：專案地圖與學習方式",
        "left": "目標：知道每課在學什麼；遇到問題先查哪份 md。\n\n觀念：課程是「帶你做」，不是只讀理論。",
        "right": "請做：\n1) 打開 docs/AI與程式入門_專案架構總覽.md 掃過標題\n2) 檔案總管進 .cursor-tutor，確認 ai_agents、etoro_trading、local_agent_demo、docs\n\n自評：一句話說明「Cursor AI」與「本機 Python 腳本」差別。",
    },
    {
        "title": "第 1 課：檔案總管 + 專案根目錄",
        "left": "目標：聽到「repo 根目錄」就知是 .cursor-tutor 那一層。\n\n觀念：路徑對了，py -m 才找得到套件。",
        "right": "請做：\n1) 進入 C:\\Users\\User\\.cursor-tutor\n2) ai_agents\\選單_啟動AI_Agent.bat（先找到即可）\n3) etoro_trading\\啟動檔\\一鍵_Ollama摘要.bat\n\n自評：誰是「統一入口」？誰是「eToro 主專案」？",
    },
    {
        "title": "第 2 課：終端機 — cd 與位置",
        "left": "目標：會用終端機站到專案根目錄。\n\n觀念：cd = change directory。",
        "right": "指令：\ncd C:\\Users\\User\\.cursor-tutor\ndir\n\n自評：進 ai_agents 要打什麼？（cd ai_agents）",
    },
    {
        "title": "第 3 課：Python — py 跑一行",
        "left": "目標：確認 Python 可執行。\n\n觀念：py -c 執行一小段字串程式。",
        "right": "指令：\npy -c \"print('hello')\"\n\n應輸出 hello。",
    },
    {
        "title": "第 4 課：py -m 與專案入口",
        "left": "目標：理解 etoro_trading 用 -m 啟動的原因。\n\n觀念：-m 用模組方式跑，路徑才正確。",
        "right": "在根目錄執行：\npy -m etoro_trading.scripts --help\n\n自評：-m 與雙擊 .bat 共通點？（正確啟動已寫好的程式）",
    },
    {
        "title": "第 5 課：pip — 套件 ≠ 付費",
        "left": "目標：會 pip install；openai「套件」≠ OpenAI「雲端計費」。\n\n觀念：連到哪個 URL 才決定要不要錢。",
        "right": "指令：\npy -m pip install -r local_agent_demo/requirements.txt\n\n閱讀：local_agent_demo/README.md 付費說明\n\n自評：為何裝 openai 套件仍可 0 元？（因 base_url 指本機 Ollama）",
    },
    {
        "title": "第 6 課：Ollama — 服務與埠",
        "left": "目標：Ollama 是背景服務；聊天視窗只是介面之一。\n\n觀念：腳本用 API 連 127.0.0.1。",
        "right": "請做：\n1) 開啟 Ollama\n2) ollama list\n3) 瀏覽器開 http://127.0.0.1:11434\n\n自評：預設埠？（11434）",
    },
    {
        "title": "第 7 課：ai_agents 選單",
        "left": "目標：入口與實際程式位置對上。\n\n觀念：.bat 先 cd 再 py -m。",
        "right": "雙擊 ai_agents\\選單_啟動AI_Agent.bat\n至少試選項 4（開 Ollama）\n選做：選項 3 local_agent_demo\n\n自評：選 2 對應哪個子命令？（ollama-summarize）",
    },
    {
        "title": "第 8 課：eToro 摘要與 log",
        "left": "目標：腳本需要輸入資料；沒 log ≠ 壞掉。\n\n觀念：daily_*.log 多來自日常自動化。",
        "right": "請做：\n1) 看 etoro_trading\\data\\automation_logs\n2) 雙擊一鍵_Ollama摘要（任一路徑）\n3) 看終端機訊息與輸出檔位置\n\n延伸：日常自動化_排程說明.md",
    },
    {
        "title": "第 9 課：讀懂小程式 config.py",
        "left": "目標：敢打開 .py；分得出「設定」與「邏輯」。\n\n觀念：MODEL 決定呼叫哪個本機模型名。",
        "right": "打開 local_agent_demo/config.py\n找：OLLAMA_BASE_URL、OLLAMA_MODEL、MAX_AGENT_STEPS\n\n自評：指一行說「這是設定」。",
    },
    {
        "title": "第 10 課：改 SYSTEM 看輸出",
        "left": "目標：提示詞影響行為；模型是機率，不保證 100%。\n\n觀念：與「寫死 if-else」不同。",
        "right": "改 agent.py 里 SYSTEM 加一句規則（如開頭加【demo】）\nOllama 開著，執行：\npy local_agent_demo\\agent.py 測試",
    },
    {
        "title": "第 11 課：.env 與本機／雲端",
        "left": "目標：設定決定流量與費用。\n\n觀念：127.0.0.1 本機；api.openai.com 雲端+金鑰。",
        "right": "閱讀 etoro_trading/.env.example 的 OLLAMA_ 列\n對照 Ollama_本機摘要設定.md\n\n自評：為何 .env 不要貼公開聊天？",
    },
    {
        "title": "第 12 課：排程與自動化心態",
        "left": "目標：自動化 = OS/排程重複你已手動成功過的指令。\n\n觀念：不是魔法，是排時間觸發。",
        "right": "閱讀 etoro_trading/docs/日常自動化_排程說明.md\n思考：每天 9 點摘要需什麼先成立？（例：Ollama、有 log）\n\n畢業：一週後可從第 7 課重跑。",
    },
    {
        "title": "學完往哪走 + 免責",
        "left": "eToro 日常 → 簡易操作流程.md\n策略+Cursor → 如何用_Cursor_當策略討論_AI_agent.md\n加工具 → local_agent_demo/tools.py\n雲端編程 → claude / codex（見 ai_agents/README）",
        "right": "與助教互動：\n「我做完第 N 課，自評：…」\n\n免責：本教材僅協助電腦操作學習；投資決策請自行判斷。",
    },
]


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for item in LESSONS:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.65))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = item["title"]
        p.font.size = PptPt(28)
        p.font.bold = True

        left_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.05), Inches(6.1), Inches(6.0))
        ltf = left_box.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = "【觀念／目標】\n" + item["left"]
        lp.font.size = PptPt(14)

        right_box = slide.shapes.add_textbox(Inches(6.75), Inches(1.05), Inches(6.1), Inches(6.0))
        rtf = right_box.text_frame
        rtf.word_wrap = True
        rp = rtf.paragraphs[0]
        rp.text = "【操作／自評】\n" + item["right"]
        rp.font.size = PptPt(14)

    prs.save(str(OUT_PPTX))
    print("Wrote", OUT_PPTX)


def build_docx() -> None:
    doc = Document()

    h = doc.add_heading("自學課程 — 雙欄介紹（Word 版）", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "左欄：觀念與目標　｜　右欄：操作步驟與自評。"
        "完整步驟仍以 Markdown 為主：docs/自學課程_AI與專案實操_一步一步.md"
    )
    p.paragraph_format.space_after = Pt(12)

    for item in LESSONS:
        doc.add_heading(item["title"], level=1)

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Cm(8.0)
        table.columns[1].width = Cm(8.0)

        c0 = table.rows[0].cells[0]
        c1 = table.rows[0].cells[1]

        p0 = c0.paragraphs[0]
        r0 = p0.add_run("觀念／目標\n\n")
        r0.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
        p0.add_run(item["left"]).font.size = Pt(10)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run("操作／自評\n\n")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
        p1.add_run(item["right"]).font.size = Pt(10)

        doc.add_paragraph()

    doc.save(str(OUT_DOCX))
    print("Wrote", OUT_DOCX)


def main() -> None:
    build_pptx()
    build_docx()


if __name__ == "__main__":
    main()
