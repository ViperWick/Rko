# -*- coding: utf-8 -*-
"""
eToro 交易程式 UI 投影片生成腳本
產生含 UI 示意圖的 PPT，方便預覽介面設計
"""
import os
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

SCRIPT_DIR = Path(__file__).resolve().parent
# 圖片來源：generate_image 產生的 assets 或本目錄 ppt_images
ASSETS_CANDIDATES = [
    SCRIPT_DIR.parent / "assets",
    Path(r"C:\Users\User\.cursor\projects\c-Users-User-cursor-tutor\assets"),
]
PPT_IMAGES_DIR = SCRIPT_DIR / "ppt_images"
# 產出放在 docs/，與其他文件統一
OUTPUT_PPT = SCRIPT_DIR.parent / "docs" / "eToro_交易程式_UI預覽.pptx"


def find_images_dir():
    """找出包含 UI 圖片的目錄（優先使用已複製的 ppt_images）"""
    if PPT_IMAGES_DIR.exists():
        for f in ["etoro_ui_strategy_overview.png", "etoro_ui_portfolio.png", "etoro_ui_trading.png"]:
            if (PPT_IMAGES_DIR / f).exists():
                return PPT_IMAGES_DIR
    for d in ASSETS_CANDIDATES:
        if d.exists():
            for f in ["etoro_ui_strategy_overview.png", "etoro_ui_portfolio.png", "etoro_ui_trading.png"]:
                if (d / f).exists():
                    return d
    return None


def ensure_ppt_images():
    """確保 ppt_images 目錄有圖片（從 assets 複製）"""
    PPT_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    for d in ASSETS_CANDIDATES:
        if not d.exists():
            continue
        for f in ["etoro_ui_strategy_overview.png", "etoro_ui_portfolio.png", "etoro_ui_trading.png"]:
            src = d / f
            dst = PPT_IMAGES_DIR / f
            if src.exists() and (not dst.exists() or src.stat().st_mtime > dst.stat().st_mtime):
                shutil.copy2(src, dst)
    return PPT_IMAGES_DIR


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide


def add_slide_with_image(prs, title, img_path, caption=""):
    """新增帶圖片的投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    if img_path and os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.1), width=Inches(8.4))
    if caption:
        capBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.6))
        capTf = capBox.text_frame
        capP = capTf.paragraphs[0]
        capP.text = caption
        capP.font.size = Pt(14)
        capP.alignment = PP_ALIGN.CENTER
    return slide


def main():
    ensure_ppt_images()
    img_dir = find_images_dir() or PPT_IMAGES_DIR

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "eToro 交易程式 UI 預覽",
        "策略總覽、投資組合、交易介面示意"
    )

    slides_info = [
        ("策略總覽", "etoro_ui_strategy_overview.png", "資產動作、賣出規則、ETF 買入清單"),
        ("投資組合", "etoro_ui_portfolio.png", "倉位表格與資產配置"),
        ("交易介面", "etoro_ui_trading.png", "買入訊號、賣出建議、操作記錄"),
    ]

    for title, filename, caption in slides_info:
        img_path = img_dir / filename
        add_slide_with_image(prs, title, str(img_path), caption)

    OUTPUT_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PPT)
    print(f"已產生 PPT：{OUTPUT_PPT}")
    return str(OUTPUT_PPT)


if __name__ == "__main__":
    main()
