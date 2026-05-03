# -*- coding: utf-8 -*-
"""
產出「量子板塊」Excel 觀察清單（多工作表、可篩選、含自填欄位）
執行: 從 repo 根目錄  py etoro_trading/scripts/export_quantum_watchlist_xlsx.py
輸出: etoro_trading/docs/量子板塊_觀察清單.xlsx
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

DOCS = Path(__file__).resolve().parent.parent / "docs"
OUT = DOCS / "量子板塊_觀察清單.xlsx"

# M1 總表（與 量子板塊_總覽與儀表板.md 同步；可自行改腳本後重跑）
# ROWS_M1_BASE：前 12 欄；ROWS_M1_結構6：財務結構標籤＋量子商業化敘事（與 docs 對照，非投資建議）
# 「可轉債／ATM」欄：指**常見股權稀釋工具**線索；**待核**=請以最新 10-Q/8-K/SEDAR 為準。
ROWS_M1_BASE = [
    ("A 純量子", "IonQ", "美國", "劍橋創新中心、韓國KISTI、ORNL、CCRM生技、Braket上架Forte Enterprise", "IONQ", "離子阱", "已上市", "是", 9425, "2025多筆ATM/大額增資公司稿；現金以最新10-Q為準；留意股本稀釋", "Tempo/#AQ、256Q願景、雙閘門忠實度與容錯長期路線", "EAR出口目的地、聯邦/軍標案資安、雲跨境資料"),
    ("A 純量子", "Rigetti", "美國", "廣達戰略投資／製造、印度C-DAC 108Q訂單、NVQLink／AFRL、Braket批量優化", "RGTI", "超導", "已上市", "是", 10655, "財報曾列高現金+投資；營收仍小、GAAP淨損常含非現金—詳10-Q", "Chiplet模組化 100→150→1000Q目標、雙閘門忠實度路線圖", "兩用量子技術、政府/DOD專案合規、技術出口文件"),
    ("A 純量子", "D-Wave", "加拿大（美國上市）", "Carahsoft公部門、Unissant聯邦、Fortune100 QCaaS、併購QCI、Southeastern聯盟", "QBTS", "退火／閘控延伸", "已上市", "是", 10881, "2025高現金與營收成長敘事、全年仍淨損；併購QCI；有息負債查10-K", "Advantage2 GA、併QCI加速閘門dual-rail、Stride混合求解", "併購整合合規、EAR、加美雙重披露"),
    ("A 純量子", "QCI", "美國", "Ciena OFC示範、併購NuCrypt／Luminar強化光子通訊", "QUBT", "光子／平台（以 IR 為準）", "已上市", "是", 14284, "2025多起巨額私募/增資稿；流動資產高、需對照營收與稀釋", "光子QKD/PQC與光通訊整合、併購擴垂直", "密碼與光設備出口分類、NIST PQC對齊"),
    ("A 純量子", "Infleqtion", "美國", "英國NQCC百位元系統、SQALE2政府補助、DOE QSC／ORNL與NVQLink", "INFQ", "中性原子", "已上市", "否", "", "SPAC+PIPE毛收益約5.5億美元級敘事；上市後燒現金率看10-Q", "Sqale大陣列、Sqorpius邏輯位元願景、高忠實度閘門", "英美政府專案、EAR、技術轉讓與NVLink跨境"),
    ("A 純量子", "Xanadu", "加拿大", "AFRL四年R&D、Corning光纖、Applied 300mm TES、A*STAR擴大合作", "XNDU", "光子", "已掛牌（新近）", "否", "", "新掛牌—以首份季報現金、研發費、燒現金為主", "Aurora室溫光子、Nature論文、損耗與容錯下一步", "AFRL國防R&D、美加出口與光電供應鏈"),
    ("A 純量子", "IQM", "芬蘭／歐盟", "芬蘭學研交機、歐盟政府專案為主（詳見IR）", "待定", "超導", "SPAC／上市進行中", "", "", "未上市—輪次/跑道見招股草稿與媒體", "150/300Q交付芬蘭、與LUMI/HPC整合", "歐盟兩用出口、FDI審查、政府採購"),
    ("A 純量子", "Pasqal", "法國／歐盟", "Azure上架、IBM／NVIDIA HPC整合、新加坡FSQS、法國CNRS", "待定", "中性原子", "規劃掛牌", "", "", "未上市—創投與政府合約；無公開季報", "中性原子上雲、可升級容錯路線圖", "歐法星多法域、雲端客戶條款、兩用物項"),
    ("A 純量子", "Quantinuum", "英國／美國", "Honeywell募資與IPO草稿、NVIDIA NVentures／Quanta等策略投資人", "待定", "離子阱／軟體", "IPO 準備", "", "", "2025大額募資與IPO草稿；合併科目見Honeywell與未來S-1", "Helios商轉、高忠實離子阱、邏輯位元與GenQAI", "跨境資料、IPO披露、軍民大客戶審查"),
    ("B 量子安全", "Arqit", "英國", "RAD電信、6WIND量子安全VPN", "ARQQ", "後量子／加密敘事", "已上市", "是", 14435, "FY25現金約数千万美元級、營收仍小；跑道依燒現金與是否增資", "NetworkSecure等量子安全連線產品、電信整合", "各國密碼進口規定、PQC對齊、企業採購資安"),
    ("B 量子安全", "01 Communique", "加拿大", "Turnium郵件加密、qLABS代幣、SuperQ整合IronCAP", "ONE", "後量子通訊", "已上市", "否", "", "TSX小盤—SEDAR查募資、認股權、燒現金", "IronCAP（NIST PQC）郵件/檔案/錢包", "PQC合規、TSX披露、代幣聯盟適用法規"),
    ("B 量子安全", "BTQ", "加拿大", "QPerfect中性原子策略投資與協定、量子簽章／演算法共研", "BTQ", "後量子加密", "已上市", "否", "", "NEO/多交易所—MD&A與私募/行使權證追蹤", "PQC+演算法與中性原子硬體共研", "PQC、兩用技術、多交易所披露"),
    ("B 量子安全", "Quantum eMotion", "加拿大", "Jmem台灣SoC、NRC IRAP、魁北克Kirq測試床", "QNC", "隨機數／加密", "TSXV", "否", "", "TSXV—補助+研發；現金與增資見季報", "QRNG SoC、與Jmem整合PQC", "FIPS140/CC若進政企、加拿大研發補助"),
    ("B 量子安全", "SEALSQ", "瑞士", "TSS美國國防PQC路線、ColibriTD QaaS、Quobly合作", "LAES", "安全晶片敘事", "已上市", "是", 14375, "2025 preliminary高現金敘事、淨損擴大；併購IC'ALPS—以審計年報為準", "QS7001等PQC晶片、Qvault TPM路線", "FIPS140-3、國防供應鏈、PQC晶片出口分類"),
    ("C 供應鏈", "SkyWater", "美國", "SQC混合量子、D-Wave等量子客戶、Trusted國防代工、IonQ收購議題（監管中）", "SKYT", "特殊代工", "已上市", "是", 10289, "營收紀錄與量子ATS成長；IonQ收購對價/條款見proxy與8-K", "Trusted晶圓、量子ATS客戶擴張", "ITAR/EAR、DMEA Trusted Foundry、CFIUS併購"),
    ("C 供應鏈", "AmpliTech", "美國", "主軸5G ORAN大額LOI；低温放大器供量子客戶評估（見年報）", "AMPG", "放大器等", "已上市", "否", "", "營收主在5G；量子利基—現金流與應收見10-Q", "低温放大器供量子感測評估", "航太國防客戶則EAR、AS9100視客戶要求"),
    ("C 供應鏈", "ASP Isotopes", "美國", "TerraPower HALEU、Isotopia醫療同位素、Renergen收購（材料非典型量子軟硬體）", "ASPI", "材料／同位素", "已上市", "是", 14374, "資本開支與融資密集—TerraPower/Renergen見8-K與貸款條款", "HALEU/同位素產能建置", "NRC核材料/出口為主、非量子專法"),
    ("D 大廠", "IBM", "美國", "Cisco聯網量子願景、BMO等Quantum Network、DOE國家量子中心", "IBM", "量子雲", "已上市", "是", 1020, "投資級財務；量子併入整體R&D/雲—見10-K分部附註", "Quantum Network、Qiskit、聯網量子長期願景", "集團出口合規、Fed採購子案"),
    ("D 大廠", "Alphabet", "美國", "Willow與英國NQCC研究資助、Early Access、Quantum Echoes等", "GOOGL", "量子研究", "已上市", "是", 6434, "現金強；量子為研發子項—10-K", "Willow、NQCC計畫、可驗證優勢研究", "集團法遵、研發出口內控"),
    ("D 大廠", "Microsoft", "美國", "Atom Computing商用量子、Azure上架Pasqal等、Majorana 1", "MSFT", "Azure Quantum", "已上市", "是", 1004, "整體財務穩健；Azure Quantum佔比小—10-Q", "Majorana、Azure硬體聚合", "集團法遵、Azure/GDPR等"),
    ("D 大廠", "Amazon", "美國", "Braket擴IonQ／Rigetti／IQM／AQT、CUDA-Q混合運算", "AMZN", "Braket", "已上市", "是", 1005, "AWS現金流強；Braket為子業務", "Braket多硬體、CUDA-Q混合", "雲責任分界、硬體供應商出口"),
    ("D 大廠", "Intel", "美國", "AIST矽量子工業化MOU、Argonne／QuTech／CQE長期研發", "INTC", "矽量子研究", "已上市", "是", 1021, "轉型期槓桿與現金見10-K；量子屬長期研發", "矽量子2030願景、AIST合作", "國際合作合規、EAR"),
    ("D 大廠", "Honeywell", "美國", "Quantinuum控股／募資／IPO敘事主要出口", "HON", "Quantinuum 關聯", "已上市", "是", 1469, "Quantinuum權益法/合併見HON附註；非單獨量子損益", "透過Quantinuum承載量子業務", "集團管制、Quantinuum IPO披露"),
]

# 每列對應 ROWS_M1_BASE 同序：(稀釋風險, 可轉債_有填是, ATM或市價配售_有填是, 燒現金程度, 商用量子營收_提要, 客戶_政府合約為主)
# 稀釋／燒現金：高、中、低、待核（相對同類標的之主觀分級，務必交叉驗證財報）
# 可轉債／ATM：是、否、待核
# 客戶：是、否、混合、待核
ROWS_M1_結構6 = [
    ("高", "待核", "是", "高", "財報多未單列「商用量子」分部；QCaaS/授權營收仍小—以10-K分部與MD&A為準", "混合"),
    ("中", "待核", "是", "高", "營收多為百萬級早期量子服務+硬體；佔比與毛利見10-Q", "混合"),
    ("中", "待核", "是", "中", "有QCaaS/授權等營收敘事且年增常披露；併購QCI後科目變動—查分部", "混合"),
    ("高", "待核", "是", "高", "多次市價/私募後與商業營收並存；光子與資安佔比請對照分部", "混合"),
    ("中", "待核", "否", "高", "上市初期；國防/政府里程碑與商業訂單認列節奏待多季10-Q驗證", "是"),
    ("待核", "待核", "待核", "待核", "新掛牌；收入多為政府R&D/補助或里程碑—首份年報前標為待核", "是"),
    ("待核", "待核", "待核", "待核", "未上市；媒體/招股敘事以歐洲學研與政府專案為主", "是"),
    ("待核", "待核", "待核", "待核", "未上市；雲+HPC訂閱與歐盟/法國政府及研究單位並存", "混合"),
    ("待核", "待核", "待核", "待核", "IPO前無公開季報；Honeywell關聯與大客戶結構待S-1", "混合"),
    ("高", "待核", "否", "高", "FY25敘事營收仍極低；產品為量子安全網路/軟體—非通用量子運算營收", "混合"),
    ("高", "待核", "待核", "高", "TSX小盤；量子安全營收體量與認列以SEDAR季報為準", "否"),
    ("高", "待核", "待核", "高", "NEO小盤；PQC/加密敘事為主—營收拆細待MD&A", "混合"),
    ("高", "待核", "待核", "高", "補助+研發期；商業QRNG與SoC營收占比見TSXV季報", "是"),
    ("中", "待核", "待核", "中", "安全晶片與服務營收為主；PQC產品線佔總營收比查年報", "是"),
    ("低", "待核", "否", "低", "晶圓代工營收為主；量子ATS為成長線—佔比見年報與投資人簡報", "混合"),
    ("中", "待核", "待核", "中", "主軸5G/ORAN大額訂單敘事；量子低温放大器多為評估/小量—佔比小", "否"),
    ("中", "待核", "待核", "高", "主業為同位素/HALEU與資本專案；與「量子運算營收」關聯弱", "混合"),
    ("低", "否", "否", "低", "量子併入整體雲與研發；無單列商用量子營收", "混合"),
    ("低", "否", "否", "低", "量子為研發無單列營收；政府研究資助屬合作非主營收來源", "混合"),
    ("低", "否", "否", "低", "Azure Quantum佔總營收極小；財報通常不單列", "混合"),
    ("低", "否", "否", "低", "Braket佔AWS營收極小", "混合"),
    ("低", "否", "否", "中", "矽量子屬長期R&D；無單列量子營收", "混合"),
    ("低", "否", "否", "低", "Quantinuum於HON附註；無單獨「商用量子」損益公開拆細", "混合"),
]

COLS_M1 = [
    "分類",
    "公司",
    "國家／區域",
    "重大合作_提要",
    "代碼",
    "技術路線",
    "上市進度",
    "eToro_API命中_20260321",
    "instrumentId",
    "財務與融資_提要",
    "研究與產品進度_提要",
    "監管與認證_提要",
    "稀釋風險_高中低待核",
    "可轉債_有填是",
    "ATM或市價配售_有填是",
    "燒現金程度_高中低待核",
    "商用量子營收_提要",
    "客戶_政府合約為主",
    "自選監控_是填是",
    "我有持倉_是填是",
    "研究筆記_自由填寫",
    "核證日期_YYYY-MM-DD",
    "下次複盤日",
]


def main():
    DOCS.mkdir(parents=True, exist_ok=True)

    if len(ROWS_M1_BASE) != len(ROWS_M1_結構6):
        raise ValueError("ROWS_M1_BASE 與 ROWS_M1_結構6 列數必須相同")
    df_m1 = pd.DataFrame(
        [
            list(a) + list(b) + ["", "", "", "", ""]
            for a, b in zip(ROWS_M1_BASE, ROWS_M1_結構6)
        ],
        columns=COLS_M1,
    )
    # instrumentId 轉成可空數字（Excel 好排序）
    df_m1["instrumentId"] = pd.to_numeric(df_m1["instrumentId"], errors="coerce")

    df_ipo = pd.DataFrame(
        {
            "標的": ["Quantinuum", "IQM", "Pasqal", "XNDU", "INFQ"],
            "國家／區域": ["英國／美國", "芬蘭／歐盟", "法國／歐盟", "加拿大", "美國"],
            "你該盯的": [
                "招股書／S-1、掛牌日、代碼",
                "SPAC 交割日、最終代碼",
                "雙掛牌時程、條款更新",
                "掛牌後首份財報、eToro 是否上架",
                "季報結構、eToro 是否上架",
            ],
            "狀態_自由填寫": [""] * 5,
            "下次檢查日": [""] * 5,
            "備註": [""] * 5,
        }
    )

    df_risk = pd.DataFrame(
        {
            "代碼": ["WIMI", "UIS", "FCCN", "ZENA", "ALMU"],
            "國家／區域": ["中國", "美國", "美國", "待查", "待查"],
            "eToro_API命中_20260321": ["是", "否", "否", "否", "未驗"],
            "instrumentId": [13659, 10737, "", "", ""],
            "說明": [
                "高風險名單常見",
                "高風險名單常見",
                "OTC 常未上架",
                "API 未命中",
                "尚未批次驗證",
            ],
            "是否納入觀察_是填是": [""] * 5,
            "筆記": [""] * 5,
        }
    )

    df_etoro_miss = pd.DataFrame(
        {
            "代碼": ["INFQ", "XNDU", "QNC", "BTQ", "ONE", "AMPG", "ZENA", "FCCN"],
            "國家／區域": ["美國", "加拿大", "加拿大", "加拿大", "加拿大", "美國", "待查", "美國"],
            "建議動作": [
                "App 搜 Infleqtion",
                "新上市可能延後上架",
                "TSXV 小板",
                "NEO 板",
                "TSX 試其他寫法",
                "API 無結果",
                "API 無結果",
                "OTC",
            ],
            "已手動確認eToro_是填是": [""] * 8,
            "實際搜尋代碼_若不同請填": [""] * 8,
        }
    )

    readme = pd.DataFrame(
        {
            "說明": [
                "本檔為研究輔助，非投資建議；損益自負。",
                "",
                "工作表「M1_總表」：主清單；含合作／財務／研究／監管提要，以及「稀釋風險」「可轉債」「ATM／市價配售」「燒現金程度」「商用量子營收_提要」「客戶_政府合約為主」（與量子板塊_財務與資本動態.md 對照；下拉僅輔助篩選，**待核**請自行查10-Q/8-K/SEDAR）。",
                "細節長文見工作表「L_參考文件與MD」所列 md（etoro_trading/docs/；非法律或投資建議）。",
                "欄位 eToro_API命中_20260321 為批次 API 驗證日；請以 App 為準不定期更新。",
                "「自選監控」「我有持倉」「研究筆記」「核證日期」「下次複盤日」皆可自由編輯。",
                "",
                "工作表「IPO_觀察」：掛牌節點追蹤。",
                "工作表「M4_高風險」：題材股，建議核年報後再決定是否觀察。",
                "工作表「eToro未命中」：待你在 App 補確認的代碼。",
                "工作表「L_參考文件與MD」：與量子板塊相關的 Markdown 清單與路徑。",
                "",
                "更新主表結構後：執行 py etoro_trading/scripts/export_quantum_watchlist_xlsx.py 可覆寫產生新檔（會覆蓋同名 xlsx）。",
                "若你已在 Excel 內填大量筆記，重跑前請先另存新檔備份。",
            ]
        }
    )

    df_docs = pd.DataFrame(
        {
            "檔名": [
                "量子板塊_總覽與儀表板.md",
                "量子板塊_重大合作與聯盟摘要.md",
                "量子板塊_財務與資本動態.md",
                "量子板塊_合作計畫與研究進度總覽.md",
                "量子板塊_監管與認證概覽.md",
            ],
            "簡述": [
                "板塊地圖、分類、與本 xlsx 對照",
                "各公司聯盟／客戶／雲上架等合作提要",
                "融資、現金跑道、併購與財報風險提要",
                "研發里程碑、產品路線、論文與示範提要",
                "出口管制、密碼／PQC、國防採購與認證提要",
            ],
            "相對路徑_etoro_trading_docs": [
                "docs/量子板塊_總覽與儀表板.md",
                "docs/量子板塊_重大合作與聯盟摘要.md",
                "docs/量子板塊_財務與資本動態.md",
                "docs/量子板塊_合作計畫與研究進度總覽.md",
                "docs/量子板塊_監管與認證概覽.md",
            ],
        }
    )

    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="0_使用說明", index=False)
        df_m1.to_excel(writer, sheet_name="M1_總表", index=False)
        df_ipo.to_excel(writer, sheet_name="IPO_觀察", index=False)
        df_risk.to_excel(writer, sheet_name="M4_高風險", index=False)
        df_etoro_miss.to_excel(writer, sheet_name="eToro未命中待查", index=False)
        df_docs.to_excel(writer, sheet_name="L_參考文件與MD", index=False)

    # 格式：凍結首列、篩選、欄寬（依整欄內容自動調整，含中文顯示寬度估算）
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    def _cell_display_units(val: object) -> float:
        """估算 Excel 欄寬用的「字元單位」：ASCII 約 1，中日韓等約 2。"""
        s = "" if val is None else str(val)
        w = 0.0
        for c in s:
            o = ord(c)
            if o < 128:
                w += 1.0
            else:
                w += 2.0
        return w

    def _autofit_sheet(ws, max_width: float = 72.0, min_width: float = 9.0) -> None:
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            best = min_width
            for row in ws.iter_rows(
                min_row=1,
                max_row=ws.max_row,
                min_col=col_idx,
                max_col=col_idx,
            ):
                for cell in row:
                    best = max(best, _cell_display_units(cell.value))
            ws.column_dimensions[letter].width = min(max(best + 1.8, min_width), max_width)

    from openpyxl.worksheet.datavalidation import DataValidation

    wb = load_workbook(OUT)
    for name in wb.sheetnames:
        ws = wb[name]
        ws.freeze_panes = "A2"
        if name == "0_使用說明":
            # 長段說明：仍給較寬單欄，但依內容上限避免過窄
            _autofit_sheet(ws, max_width=100.0, min_width=12.0)
            continue
        ws.auto_filter.ref = ws.dimensions
        mw = 90.0 if name == "M1_總表" else 72.0
        _autofit_sheet(ws, max_width=mw, min_width=9.0)

    # M1：可篩選欄位下拉（類似勾選；可改儲存格覆寫時需允許或關閉驗證）
    ws_m1 = wb["M1_總表"]
    last_row = max(2, ws_m1.max_row)
    _m1_dvs = [
        ("M", '"高,中,低,待核"'),
        ("N", '"是,否,待核"'),
        ("O", '"是,否,待核"'),
        ("P", '"高,中,低,待核"'),
        ("R", '"是,否,混合,待核"'),
    ]
    for col_letter, formula in _m1_dvs:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws_m1.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{last_row}")

    wb.save(OUT)
    print(f"已寫入: {OUT}")


if __name__ == "__main__":
    main()
