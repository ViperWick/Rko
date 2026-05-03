# -*- coding: utf-8 -*-
"""
AI Agent 完整架構與流程 PPT 生成腳本
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# 腳本所在目錄，用於尋找圖片
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
IMAGES_DIR = os.path.join(SCRIPT_DIR, "images")

def add_title_slide(prs, title, subtitle=""):
    """新增標題投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide

def add_content_slide(prs, title, content_list, bullet=True):
    """新增內容投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, item in enumerate(content_list):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
    return slide

def add_section_slide(prs, title):
    """新增章節投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return slide

def add_table_slide(prs, title, headers, rows):
    """新增表格投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    # 簡單用文字呈現表格
    body = slide.placeholders[1].text_frame
    body.clear()
    header_text = "  |  ".join(headers)
    p = body.paragraphs[0]
    p.text = header_text
    p.font.bold = True
    p.font.size = Pt(14)
    for row in rows:
        p = body.add_paragraph()
        p.text = "  |  ".join(str(c) for c in row)
        p.font.size = Pt(12)
    return slide

def add_slide_with_image(prs, title, image_filename, caption=""):
    """新增帶流程圖的投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白
    # 標題
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    # 圖片（置中，寬 7 吋）
    img_path = os.path.join(IMAGES_DIR, image_filename)
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.2), width=Inches(7))
    # 說明文字
    if caption:
        capBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(0.8))
        capTf = capBox.text_frame
        capP = capTf.paragraphs[0]
        capP.text = caption
        capP.font.size = Pt(14)
        capP.alignment = PP_ALIGN.CENTER
    return slide

def add_slide_with_two_images(prs, title, img1_filename, img2_filename, caption1="", caption2=""):
    """新增帶兩張流程圖的投影片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    # 兩張圖並排
    path1 = os.path.join(IMAGES_DIR, img1_filename)
    path2 = os.path.join(IMAGES_DIR, img2_filename)
    if os.path.exists(path1):
        slide.shapes.add_picture(path1, Inches(0.5), Inches(1), width=Inches(4.2))
    if os.path.exists(path2):
        slide.shapes.add_picture(path2, Inches(5.2), Inches(1), width=Inches(4.2))
    if caption1 or caption2:
        capBox = slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(9), Inches(1))
        capTf = capBox.text_frame
        capP = capTf.paragraphs[0]
        capP.text = f"{caption1}  |  {caption2}" if caption1 and caption2 else (caption1 or caption2)
        capP.font.size = Pt(12)
        capP.alignment = PP_ALIGN.CENTER
    return slide

def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ===== 封面 =====
    add_title_slide(prs, 
        "AI Agent 完整架構與實作指南",
        "從概念到 Slack、GitHub、股票 API 整合\n完整流程與架構說明")

    # ===== Part 1: 什麼是 Agent =====
    add_section_slide(prs, "Part 1：什麼是 AI Agent？")

    add_content_slide(prs, "Agent 的定義", [
        "AI Agent = 能自主感知、決策、行動的系統",
        "不是單純聊天機器人，而是會「做事」的 AI",
        "",
        "傳統 AI vs Agent：",
        "• 傳統：回答問題、需要你一步步下指令",
        "• Agent：自己規劃步驟、調用工具、完成任務",
        "",
        "核心公式：Agent = LLM（大腦）+ 工具（手腳）+ 記憶 + 規劃迴圈"
    ])

    add_content_slide(prs, "Agent 調動的是什麼？", [
        "調動的是「工具」（Tools / Function Calling）",
        "",
        "流程示意：",
        "1. 用戶：「幫我查台北明天天氣」",
        "2. LLM 判斷：需要調用「天氣 API」",
        "3. 輸出結構化請求：{ tool: get_weather, params: {...} }",
        "4. 系統執行 API，取得結果",
        "5. 結果回傳 LLM，整理成自然語言回覆",
        "",
        "重點：LLM 只「決定」調用什麼，實際執行由系統完成"
    ])

    add_slide_with_image(prs, "工具調用流程圖", "tool_calling_flow.png",
        "用戶提問 → LLM 決策 → 調用工具 → 取得結果 → 回覆用戶")

    add_content_slide(prs, "常見工具類型", [
        "搜尋：網頁搜尋、維基百科、學術論文",
        "資料：資料庫查詢、API 呼叫",
        "執行：執行代碼、跑腳本、終端機指令",
        "檔案：讀寫檔案、管理專案",
        "通訊：Slack、Discord、Email",
        "應用：操作瀏覽器、控制 IDE、券商下單 API"
    ])

    # ===== Part 2: 架構 =====
    add_section_slide(prs, "Part 2：Agent 四層架構")

    add_content_slide(prs, "四層核心架構", [
        "第一層 - LLM（大腦）：推理、決策、規劃",
        "第二層 - 工具（手臂）：API、搜尋、執行、檔案操作",
        "第三層 - 記憶：短期（當前對話）、長期（向量庫、歷史）",
        "第四層 - 反饋迴圈：感知 → 推理 → 行動 → 觀察 → 評估 → 改進",
        "",
        "代理迴圈：ReAct 模式",
        "推理(Reason) → 行動(Act) → 觀察(Observe) → 再推理..."
    ])

    add_slide_with_two_images(prs, "四層架構與 ReAct 迴圈流程圖",
        "agent_architecture.png", "react_loop.png",
        "四層架構", "ReAct 迴圈")

    add_content_slide(prs, "AI 使用方式光譜", [
        "被動回應 ←———————————————→ 主動執行",
        "",
        "純聊天（ChatGPT）→ 增強聊天（帶工具）→ 單一 Agent（Cursor）→ 多 Agent 協作",
        "",
        "• 純聊天：問答、翻譯、摘要、寫作",
        "• 增強：搜尋、計算、查資料、畫圖",
        "• 單一 Agent：寫代碼、改檔案、執行指令、多步驟任務",
        "• 多 Agent：分工合作、自動化、雲端代理"
    ])

    # ===== Part 3: 整合實作 =====
    add_section_slide(prs, "Part 3：如何整合 Slack、GitHub、API？")

    add_content_slide(prs, "基本原理：API + 授權", [
        "讓 AI 操作外部服務 = 調用該服務的 API",
        "",
        "授權方式：",
        "• API Key：一串金鑰，放在請求裡",
        "• OAuth：用戶授權登入，取得 token（較安全）",
        "• Webhook：服務主動推事件給你",
        "",
        "流程：AI Agent → 調用 API（帶授權）→ 服務端執行 → 回傳結果"
    ])

    add_content_slide(prs, "三種實作路線", [
        "路線 A - 自動化平台（最簡單）：",
        "  Zapier、n8n、Make：視覺化拖拉，6000+ 應用串接",
        "",
        "路線 B - 自己寫 Agent：",
        "  LangChain、LlamaIndex + Slack/GitHub SDK 當工具",
        "",
        "路線 C - 現成整合：",
        "  GitHub Copilot + Slack、ArcadeAI SlackAgent、Eliza"
    ])

    add_slide_with_image(prs, "服務整合架構圖", "integration_architecture.png",
        "AI Agent 透過 OAuth/API Key 串接 Slack、GitHub、Gmail、券商等服務")

    add_content_slide(prs, "常見服務 API 支援", [
        "Slack：發訊息、讀頻道、管理頻道、搜尋",
        "GitHub：開 Issue、PR、管理 Repo、Actions",
        "Gmail：讀信、寄信、標籤",
        "Google Calendar：查行程、建立事件",
        "Notion：讀寫頁面、資料庫",
        "股票/券商：行情查詢、下單（富途、IB、富果等）"
    ])

    # ===== Part 4: 股票 API =====
    add_section_slide(prs, "Part 4：股票 API 與自動下單")

    add_content_slide(prs, "技術上可行，但風險高", [
        "券商 API 範例：",
        "• 富途 OpenAPI：港股、美股、日股，Python SDK",
        "• Interactive Brokers：全球市場",
        "• 富果 Fugle：台股（留意官方公告）",
        "",
        "架構：AI 分析 → 調用下單 API → 券商執行",
        "",
        "⚠️ 務必：模擬帳戶測試、小額實測、設風控、人工覆核"
    ])

    add_content_slide(prs, "風險與注意事項", [
        "技術風險：程式 bug、網路延遲、API 異常",
        "市場風險：價格劇烈波動、流動性不足",
        "策略風險：止損設錯、邏輯錯誤",
        "安全風險：API Key 外洩",
        "",
        "建議：模擬環境 → 小額 → 風控（單筆上限、總曝險）→ 人工覆核"
    ])

    # ===== Part 5: 完整流程圖 =====
    add_section_slide(prs, "Part 5：完整架構總覽")

    add_content_slide(prs, "Agent 可調動的服務架構", [
        "通訊：Slack、Discord",
        "程式碼：GitHub、GitLab",
        "辦公：Gmail、日曆、Notion",
        "金融：券商 API（行情、下單）",
        "",
        "共通：OAuth / API Key 授權 → AI Agent（LLM + Tools）",
        "",
        "選擇：簡單用 Zapier/n8n；進階自己寫 Agent"
    ])

    add_slide_with_image(prs, "完整架構總覽圖", "integration_architecture.png",
        "所有服務透過授權與 AI Agent 串接，形成完整自動化生態")

    add_content_slide(prs, "實作建議總覽", [
        "Slack + GitHub 自動化 → Zapier / n8n",
        "查股票、新聞、報價 → 用行情 API 當工具（不下單）",
        "自動下單 → 先模擬、小額、人工覆核、了解法規",
        "",
        "安全：API Key 勿寫進程式碼、勿上傳公開 repo"
    ])

    # ===== 結尾 =====
    add_title_slide(prs, 
        "謝謝",
        "AI Agent 時代，理解架構才能善用工具\n有任何問題歡迎繼續討論")

    # 儲存（固定輸出到腳本所在目錄，使用英文檔名避免編碼問題）
    output_path = os.path.join(SCRIPT_DIR, "AI_Agent_Guide.pptx")
    prs.save(output_path)
    print(f"PPT 已生成：{output_path}")
    return output_path

if __name__ == "__main__":
    main()
