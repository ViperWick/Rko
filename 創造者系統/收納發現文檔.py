# -*- coding: utf-8 -*-
"""
發現文檔收納腳本
掃描「發現文檔收納」資料夾內的 .txt / .docx / .pptx / .xlsx，
產生/更新「發現文檔索引.md」。
"""

from pathlib import Path
from datetime import datetime
from typing import List, Tuple

# 設定：收納資料夾名、索引檔名、支援的副檔名
SCRIPT_DIR = Path(__file__).resolve().parent
INBOX_DIR = SCRIPT_DIR / "發現文檔收納"
INDEX_FILE = SCRIPT_DIR / "發現文檔索引.md"
SUPPORTED_EXT = {".txt", ".docx", ".pptx", ".xlsx"}
PREVIEW_MAX = 120  # 預覽字數


def get_preview_txt(path: Path) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read(PREVIEW_MAX * 2)
        text = "".join(raw.split())[:PREVIEW_MAX]
        return (text + "…") if len(text) >= PREVIEW_MAX else text or "—"
    except Exception:
        return "—"


def get_preview_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(path)
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
                if sum(len(s) for s in parts) >= PREVIEW_MAX:
                    break
        text = "".join(parts)[:PREVIEW_MAX]
        return (text + "…") if len(text) >= PREVIEW_MAX else text or "—"
    except ImportError:
        return "（需安裝 python-docx）"
    except Exception:
        return "—"


def get_preview_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
                    if sum(len(s) for s in parts) >= PREVIEW_MAX:
                        break
            if sum(len(s) for s in parts) >= PREVIEW_MAX:
                break
        text = "".join(parts)[:PREVIEW_MAX]
        return (text + "…") if len(text) >= PREVIEW_MAX else text or "—"
    except ImportError:
        return "（需安裝 python-pptx）"
    except Exception:
        return "—"


def get_preview_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        if not ws:
            return "—"
        parts = []
        for row in ws.iter_rows(max_row=10, max_col=3, values_only=True):
            for c in row:
                if c is not None and str(c).strip():
                    parts.append(str(c).strip())
                    if sum(len(s) for s in parts) >= PREVIEW_MAX:
                        break
            if sum(len(s) for s in parts) >= PREVIEW_MAX:
                break
        wb.close()
        text = "".join(parts)[:PREVIEW_MAX]
        return (text + "…") if len(text) >= PREVIEW_MAX else text or "—"
    except ImportError:
        return "（需安裝 openpyxl）"
    except Exception:
        return "—"


def get_preview(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return get_preview_txt(path)
    if ext == ".docx":
        return get_preview_docx(path)
    if ext == ".pptx":
        return get_preview_pptx(path)
    if ext == ".xlsx":
        return get_preview_xlsx(path)
    return "—"


def escape_cell(s: str) -> str:
    """避免 Markdown 表格斷列。"""
    if not s:
        return ""
    return s.replace("|", "｜").replace("\n", " ").strip()


def collect_files() -> List[Tuple[Path, datetime]]:
    """收集收納資料夾內所有支援的檔案（含子資料夾），回傳 (path, mtime)。"""
    if not INBOX_DIR.is_dir():
        return []
    out = []
    for p in INBOX_DIR.rglob("*"):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
            try:
                mtime = datetime.fromtimestamp(p.stat().st_mtime)
                out.append((p, mtime))
            except OSError:
                pass
    return sorted(out, key=lambda x: x[1], reverse=True)


def main():
    INBOX_DIR.mkdir(parents=True, exist_ok=True)
    rows = collect_files()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    lines = [
        "# 發現文檔索引",
        "",
        "> 由「收納發現文檔」腳本自動產生。把 Word / PPT / Excel / TXT 丟進 `發現文檔收納/` 後執行腳本即可更新此索引。",
        "",
        f"**最後更新**：{now_str}",
        "",
        "---",
        "",
        "## 已收納檔案",
        "",
        "| 檔名 | 類型 | 修改日期 | 預覽 |",
        "|------|------|----------|------|",
    ]

    for path, mtime in rows:
        rel = path.relative_to(INBOX_DIR)
        name = rel.name
        ext = path.suffix.lower()
        type_name = {"": "TXT", ".txt": "TXT", ".docx": "Word", ".pptx": "PPT", ".xlsx": "Excel"}.get(ext, ext[1:].upper())
        date_str = mtime.strftime("%Y-%m-%d %H:%M")
        preview = escape_cell(get_preview(path))
        # 相對路徑方便在索引同目錄下點開（部分編輯器可點）
        rel_link = f"發現文檔收納/{rel.as_posix()}"
        lines.append(f"| [{name}]({rel_link}) | {type_name} | {date_str} | {preview} |")

    if not rows:
        lines.append("| （尚無檔案） | — | — | 請將 .docx / .pptx / .xlsx / .txt 放入「發現文檔收納」後再執行 `收納發現文檔.py` |")

    lines.extend(["", "---", "", "*收納資料夾路徑：`創造者系統/發現文檔收納/`*", ""])
    INDEX_FILE.write_text("\n".join(lines), encoding="utf-8")
    print(f"已收納 {len(rows)} 個檔案，索引已寫入：{INDEX_FILE.name}")


if __name__ == "__main__":
    main()
